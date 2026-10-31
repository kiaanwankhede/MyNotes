#!/usr/bin/env python3
"""Builds MASTER_COVERAGE.xlsx from build/master_rows.json (measured) plus recorded rows (Craft, sites, other repos)
and build/master_overrides.json (manual updates). Run after build/master_analyze.py.
Usage: python3 -I build/master_xlsx.py"""
import json, os, re, datetime, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TODAY=datetime.date.today().isoformat()
rows=json.load(open(f'{ROOT}/build/master_rows.json'))
ov=json.load(open(f'{ROOT}/build/master_overrides.json')) if os.path.exists(f'{ROOT}/build/master_overrides.json') else {}

def add(**kw):
    base=dict(source='',group='',item='',path='',concept='',read_status='',coverage='',score=None,where='',missing='',linked='',topic='',priority='Medium',note='',method='Recorded: from build notes')
    base.update(kw); rows.append(base)
for r in rows: r.setdefault('method','Measured: shared terms with the books')

by_base={}
for r in rows:
    if r['source']=='Course repo file': by_base.setdefault(r['item'].lower(),r)

# ---- Craft index pages ----
craft_pages=[('Best Practice','Foundations home ("How to study this course")','Covered'),('Class 1 - Python','Foundations f1','Covered'),('Class 2 - Windows, Python Refresher and Pydantic','Foundations f1','Covered'),('Class 3 & 4 - Pydantic & AI & Agentic AI Basics','Foundations f2, Pydantic home','Covered'),('11th July - Python and Agents','Foundations f4','Covered'),('12th July - Agents, Python and Langchain','Foundations f4','Covered'),('18th-19th July - Langchain','LangChain ch2, ch3','Covered'),('25th - 26th July - Langchain Continue','LangChain ch4, ch5','Covered'),('1st Aug','LangChain ch6','Covered'),('8-9th Aug','(page is empty on Craft)','No teaching content'),('16th Aug','LangChain ch9','Covered'),('Job Opportunity','(recruiting; skipped on purpose)','No teaching content'),('Udemy','(sponsor links; skipped on purpose)','No teaching content'),('22 Aug - Runtime + HITL','LangChain ch10','Covered'),('MCP Starts','LangChain ch11 to ch17','Covered'),('Project - AI News Letter','LangChain ch18 (newsletter project)','Covered'),('COSS Job opportunity','(recruiting; skipped on purpose)','No teaching content')]
for t,w,c in craft_pages:
    add(source='Craft page',group='Index pages',item=t,path='bugs-sleep-6uj.craft.me/agentic3',concept=t,read_status='Read in full',coverage=c,where=w,topic='Craft',priority='Medium' if c=='Covered' else 'Low',note='')

# ---- Craft attachments ----
def proxy(name):
    r=by_base.get(name.lower()); return r
att=[
('AI_Industry_Handbook_2026_Latest.pdf','Home','PDF',None,'Skipped on purpose','No teaching content','Business handout marked not for redistribution'),
('Class-1---PythonAI.pdf','Class 1','PDF','Class-1-Python-and-AI.pdf','Text extracted',None,''),
('Agentic-3-0.pdf','Class 2','PDF','Agentic-3-0.pdf','Text extracted',None,''),
('Class-3-4.pdf','Classes 3 and 4','PDF','Class-3-4.pdf','Text extracted',None,''),
('Untitled-scene (3).pdf','12 July','PDF','Untitled-scene.pdf','Text extracted',None,''),
('Agentic-Python.pdf','12 July','PDF','Agentic-Python.pdf','Text extracted',None,''),
('Langchain.pdf (22 MB)','18-19 July','PDF','Langchain.pdf','Text extracted',None,''),
('Langchain (2).pdf (24 MB)','25-26 July','PDF','Langchain.pdf','Same name family as the repo copy; byte equality not checked',None,''),
('Langchain (1).pdf (24 MB)','25-26 July','PDF','Langchain.pdf','Same name family as the repo copy; byte equality not checked',None,''),
('Langchain-2 (1).pdf','25-26 July','PDF','Langchain-2.pdf','Text extracted',None,''),
('Langchain-2.pdf','25-26 July','PDF','Langchain-2.pdf','Text extracted',None,''),
('Langchain-2 (2).pdf','1 Aug','PDF','Langchain-2.pdf','Same name family as the repo copy; byte equality not checked',None,''),
('Runtime-HITL.pdf','22 Aug','PDF','Runtime-HITL.pdf','Text extracted',None,''),
('MCP-Lifecycle.pdf (1.4 MB)','MCP Starts','PDF','MCP-Lifecycle.pdf','Text extracted',None,''),
('MCP-Lifecycle.pdf (1.8 MB, second version)','MCP Starts','PDF',None,'Read 31 Oct','Covered','LangChain ch11 section "What the Craft MCP whiteboards add"'),
('Complete-MCP.pdf','MCP Starts','PDF','Complete-MCP.pdf','Text extracted',None,''),
('MCP-Primitives.pdf','MCP Starts','PDF','MCP-Primitives.pdf','Text extracted',None,''),
('MCP-Server.pdf','MCP Starts','PDF',None,'Read 31 Oct','Covered','LangChain ch11 section "What the Craft MCP whiteboards add"'),
('MCP-Client.pdf','MCP Starts','PDF',None,'Read 31 Oct','Covered','LangChain ch11 section "What the Craft MCP whiteboards add"'),
('AI-Updater-Project.pdf','MCP Starts','PDF','AI-Updater-Project.pdf','Text extracted',None,''),
('Xunova-JD-Senior-AI-Engineer.pdf','Job Opportunity','PDF',None,'Skipped on purpose','No teaching content','Job ad'),
('Xunova-JD-Senior-Data-Scientist.pdf','Job Opportunity','PDF',None,'Skipped on purpose','No teaching content','Job ad'),
('Revision_Parts_1-4.ipynb','25-26 July','Notebook','Revision_Parts_1-4.ipynb','Read from the repo copy',None,''),
('Assignment_2_Agents_and_Prebuilt_Middleware.ipynb','16 Aug','Notebook','Assignment_2_Agents_and_Prebuilt_Middleware.ipynb','Read from the repo copy',None,''),
('Assignment_LangChain_Fundamentals_SOLUTIONS.ipynb','16 Aug','Notebook','Assignment_LangChain_Fundamentals_SOLUTIONS.ipynb','Read from the repo copy',None,''),
('Runtime_and_HITL.ipynb','22 Aug','Notebook','Runtime_and_HITL.ipynb','Read from the repo copy',None,''),
('Runtime-HITL.excalidraw','22 Aug','Diagram file',None,'NOT OPENED (its PDF export was read)','Covered','Same content as Runtime-HITL.pdf, which is covered in chapter 10'),
('Colab: Python Refresher and Pydantic','Class 2','Colab',None,'Opened; compared with the repo notebook','Covered','Python only; Foundations f1'),
('Colab: Pydantic From Scratch','Classes 3 and 4','Colab','Pydantic_From_Scratch.ipynb','Opened; compared by title and opening cells',None,''),
('Colab: Python and Agents','11 July','Colab',None,'Opened and read','Covered','Foundations f4'),
('Colab: verbose tool agent','12 July','Colab',None,'Opened and read (contains a key; not copied)','Covered','Foundations f4'),
('Colab: Langchain Collab (models, streaming)','18-19 July','Colab',None,'Opened; compared with the repo notebook (contains keys; not copied)','Covered','LangChain ch3'),
('Colab: Langchain_Continue (CineBot)','25-26 July','Colab',None,'Opened; compared with the repo notebook','Covered','LangChain ch4'),
('Colab: Student Reference, Structured Output and Tools','1 Aug','Colab','Student_Reference_Structured_Output_and_Tools.ipynb','Opened; compared with the repo notebook',None,''),
('Colab: Custom Middleware','16 Aug','Colab',None,'Opened and read in full (no repo copy)','Covered','LangChain ch9 section "The live-class Colab"'),
('Colab: Runtime and HITL','22 Aug','Colab','Runtime_and_HITL.ipynb','Opened; compared with the repo notebook',None,''),
('Colab: MCP adapter notebook (Untitled2)','MCP Starts','Colab',None,'Opened and read in full (no repo copy)','Covered','LangChain ch18 section "The MCP adapter Colab"'),
('Video: Class 1 setup (YouTube)','Class 1','Video',None,'NOT WATCHED','Not covered','Transcript not obtained'),
('Video: Class 1 second link (YouTube)','Class 1','Video',None,'NOT WATCHED','Not covered','Transcript not obtained'),
('Video: Class 2 (2:45 min)','Class 2','Video',None,'NOT WATCHED','Not covered','Transcript not obtained'),
('Video: Classes 3 and 4 (23:42 min)','Classes 3 and 4','Video',None,'NOT WATCHED','Not covered','Transcript not obtained'),
('Video: 25-26 July (20:12 min)','25-26 July','Video',None,'NOT WATCHED','Not covered','Transcript not obtained'),
]
for name,entry,kind,repo_copy,status,cov,note in att:
    c=cov; sc=None; wh=''; meth='Recorded: from build notes'
    if repo_copy:
        pr=proxy(repo_copy)
        if pr: c=pr['coverage']; sc=pr['score']; wh=pr['where']; meth='Proxy: the repo copy was measured'; note=(note+'; ' if note else '')+f'measured on repo file {repo_copy}'
        else: c='Not measured'
    add(source='Craft attachment',group=kind,item=name,path=f'Craft entry: {entry}',concept=name,read_status=status,coverage=c,score=sc,where=wh or note,topic='Craft',priority='Medium' if c not in ('No teaching content',) else 'Low',note=note,method=meth)

# ---- Craft links ----
for t,s_,c,n in [('Mayank Doubt Solver (ChatGPT link)','Not opened','Not covered','A ChatGPT tutor; cannot be read as a document'),('github.com/x1xhlol/system-prompts-and-models-of-ai-tools','Not opened','Not covered','Collected system prompts of AI tools; optional'),('MCP Inspector link with a token in the URL','Not opened','No teaching content','Local link; the token is not copied'),('Udemy and Mesh API and Kimi links','Skipped on purpose','No teaching content','Sponsor links')]:
    add(source='Craft page',group='Other links',item=t,path='Craft page',concept=t,read_status=s_,coverage=c,topic='Craft',priority='Low',note=n)

# ---- Instructor sites ----
sites=[('pydantic-with-mayank.netlify.app','Read in full','Covered','Pydantic p1 to p12'),('ai-terms-with-mayank.netlify.app','Read in full','Covered','Foundations f2'),('context-with-mayank.netlify.app','Read in full','Covered','Foundations f2'),('python-agents-with-mayank.netlify.app','Read in full','Covered','Foundations f4'),('mcp-lifecycle.netlify.app','Read in full','Covered','LangChain ch12 to ch14'),('mcp-lifecycle-simulator.netlify.app','Read in full','Covered','LangChain ch12, ch13'),('mcp-legacy-vs-modern.netlify.app','Read in full','Covered','LangChain ch17'),('ai-automation-with-mayank.netlify.app: guide APIs','Read in full','Covered','Foundations f1A'),('ai-automation-with-mayank.netlify.app: guide economics (subscription vs API)','Read in full','Covered','Foundations f2'),('ai-automation-with-mayank.netlify.app: guide open source','Read in full','Covered','Foundations f2'),('ai-automation-with-mayank.netlify.app: guide MCP','Read in full','Covered','LangChain ch11'),('ai-automation-with-mayank.netlify.app: guide RAG','Read in full','Covered','RAG r1'),('ai-automation-with-mayank.netlify.app: guide production readiness','Read in full','Covered','LangChain ch8'),('ai-automation-with-mayank.netlify.app: guide multi-agent','Read in full','Covered','LangChain ch19'),('mini-contra.netlify.app','Opened','No teaching content','A small arcade-style game'),('resilient-marzipan-839ad0.netlify.app','Not opened','No teaching content','Listed as "will not work" on Craft')]
for t,s_,c,w in sites:
    add(source='Instructor site',group='Netlify',item=t,path=t.split(':')[0],concept=t,read_status=s_,coverage=c,where=w,topic='Sites',priority='Medium' if c=='Covered' else 'Low')

# ---- Other repos ----
add(source='Other repo',group='GitHub',item='mayank953/meridian-ai-learner',path='github.com/mayank953/meridian-ai-learner',concept='The Meridian project code (same files as the Weekend 14 folder)',read_status='README, code and frontend digest read; 17 tests run',coverage='Covered',where='LangChain ch20, RAG r1 to r5',topic='Meridian / RAG',note='Files themselves are rows under Course repo file, Weekend 14')
add(source='Other repo',group='GitHub',item='patoles/agent-flow',path='github.com/patoles/agent-flow',concept='An agent-flow visualiser listed on the Craft page',read_status='README and folder layout only',coverage='Not covered',where='',topic='LangChain',priority='Low',note='Source code not read')

# ---- derive action + status, apply overrides ----
def location(r):
    p=r['path'].lower(); src=r['source']
    if src=='Pydantic docs':
        m=[('models','Pydantic ch2'),('fields','ch3'),('alias','ch3'),('types','ch3'),('validators','ch4'),('validation_decorator','ch4'),('serialization','ch5'),('unions','ch6'),('dataclasses','ch6'),('type_adapter','ch6'),('forward_annotations','ch6'),('strict_mode','ch7'),('json_schema','ch7'),('json.md','ch7'),('settings','ch8'),('errors/','ch9'),('conversion_table','ch2 or ch7'),('experimental','ch12'),('performance','ch12'),('config','ch2'),('examples/','ch12'),('integrations/','ch12'),('internals/','ch12'),('migration','Pydantic home'),('install','Foundations f1 / Pydantic home')]
        for k,v in m:
            if k in p: return v if v.startswith('Pydantic') else 'Pydantic '+v
        return 'Pydantic home'
    if src=='LangChain docs':
        m=[('mcp','LangChain ch18'),('multi-agent','ch19'),('structured-output','ch4'),('test/','ch8'),('voice-agent','ch10a'),('streaming','ch10a'),('human-in-the-loop','ch10'),('runtime','ch10'),('middleware','ch7 to ch9'),('guardrails','ch8'),('context-engineering','ch7'),('short-term','ch6'),('long-term','ch5 or ch6'),('tools','ch5'),('messages','ch3'),('models','ch3'),('agents','ch1 or ch6'),('concepts/','ch2'),('migrate','ch2'),('releases','ch2'),('common-errors','ch2 or glossary'),('versioning','ch2')]
        for k,v in m:
            if k in p: return v if v.startswith('LangChain') else 'LangChain '+v
        return 'LangChain ch2'
    if src=='MCP docs':
        m=[('inspector','MCP book: new chapter "Inspector and debugging"'),('debugging','MCP book: new chapter "Inspector and debugging"'),('security','MCP book: new chapter "Security and authorization"'),('authorization','MCP book: new chapter "Security and authorization"'),('transports','MCP ch14'),('build-server','MCP ch14'),('build-client','MCP ch16'),('connect-','MCP ch14 or ch16'),('architecture','MCP ch11'),('server-concepts','MCP ch12'),('client-concepts','MCP ch12'),('learn/versioning','MCP ch17'),('versioning','MCP ch17'),('changelog','MCP ch17'),('deprecated','MCP ch17'),('elicitation','MCP ch13'),('sampling','MCP ch17'),('roots','MCP ch13'),('server/','MCP ch12'),('patterns','MCP ch13 or ch17'),('utilities','MCP ch12'),('schema','MCP ch12'),('extensions','MCP: optional chapter'),('community','Not planned'),('roadmap','MCP ch17')]
        for k,v in m:
            if k in p: return v
        return 'MCP ch11'
    return {'MCP':'MCP chapters','Pydantic':'Pydantic book','LangChain':'LangChain book','Meridian / RAG':'RAG book or LangChain ch20','Class notes':'matching chapter','Foundations':'Foundations f1 or f4'}.get(r.get('topic',''),'decide')
n=collections.Counter()
for i,r in enumerate(rows):
    pre={'Course repo file':'R','LangChain docs':'L','MCP docs':'M','Pydantic docs':'P','Meridian slides':'S','Craft page':'C','Craft attachment':'C','Instructor site':'N','Other repo':'G'}[r['source']]
    n[pre]+=1; r['id']=f'{pre}-{n[pre]:03d}'
    key=r['path']
    o=ov.get(key) or ov.get(r['item'])
    if o:
        for k,v in o.items(): r[k]=v
    cov=r['coverage']
    if 'action' not in r or not r.get('action'):
        loc_=location(r)
        if cov=='Covered': r['action']='None'; r['status']='Done'
        elif cov=='Partly covered': r['action']=f'Add the missing concepts to {loc_}'; r['status']='To do' if r['priority'] in ('High','Medium') else 'Needs your decision'
        elif cov=='Not covered':
            if r['priority'] in ('High','Medium'): r['action']=f'Read and add to {loc_}'; r['status']='To do'
            else: r['action']='Decide whether to cover at all'; r['status']='Needs your decision'
        elif cov=='Not measured': r['action']='Check by hand'; r['status']='To do'
        else: r['action']='None'; r['status']='Not needed'
    if o and 'status' not in o and 'action' in o: r['status']=o.get('status','To do')

# ---------------- workbook ----------------
wb=Workbook(); ws=wb.active; ws.title='MASTER'
FONT='Arial'
hdr=['ID','Source','Group','Item','Path or URL','What it teaches','Read status','Concept covered?','Overlap score','Method','Where covered (book:chapter and matches)','Missing concepts (evidence)','Priority','Action needed','Action status','Last checked','Notes']
ws.append(hdr)
for r in rows:
    ws.append([r['id'],r['source'],r['group'],r['item'],r['path'],(r['concept'] or '')[:300],r['read_status'],r['coverage'],r['score'],r['method'],r['where'],r['missing'],r['priority'],r['action'],r['status'],r.get('checked',TODAY),r['note']])
widths=[8,17,22,34,48,52,20,18,9,26,40,44,9,40,16,12,40]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
hfill=PatternFill('solid',fgColor='1F3864'); thin=Side(style='thin',color='D9D9D9')
for c in ws[1]:
    c.font=Font(name=FONT,bold=True,color='FFFFFF',size=10); c.fill=hfill; c.alignment=Alignment(wrap_text=True,vertical='center')
ws.row_dimensions[1].height=32
last=ws.max_row
for row in ws.iter_rows(min_row=2,max_row=last):
    for c in row:
        c.font=Font(name=FONT,size=9); c.alignment=Alignment(wrap_text=(c.column in (4,5,6,11,12,14,17)),vertical='top')
    row[8].number_format='0%'
ws.freeze_panes='E2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(hdr))}{last}'
rng=f'H2:H{last}'
for text,color in [('Covered','C6EFCE'),('Partly covered','FFEB9C'),('Not covered','FFC7CE'),('Not measured','FFC7CE'),('Excluded by you','D9D9D9'),('No teaching content','EDEDED'),('Not applicable','EDEDED')]:
    ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=[f'"{text}"'],fill=PatternFill('solid',bgColor=color,fgColor=color)))
dv=DataValidation(type='list',formula1='"Done,To do,Needs your decision,Not needed,Skipped on purpose"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'O2:O{last}')
dv2=DataValidation(type='list',formula1='"Covered,Partly covered,Not covered,Not measured,Excluded by you,No teaching content,Not applicable"',allow_blank=True); ws.add_data_validation(dv2); dv2.add(rng)

# ---- Summary (live formulas) ----
sm=wb.create_sheet('SUMMARY')
sources=['Course repo file','LangChain docs','MCP docs','Pydantic docs','Meridian slides','Craft page','Craft attachment','Instructor site','Other repo']
cats=['Covered','Partly covered','Not covered','Not measured','Excluded by you','No teaching content','Not applicable']
sm['A1']='Master coverage summary'; sm['A1'].font=Font(name=FONT,bold=True,size=14)
sm['A2']=f'Counts update from the MASTER sheet. Generated {TODAY}. "Concept covered?" is measured by comparing each item\'s distinctive terms with the book text (see LEGEND).'; sm['A2'].font=Font(name=FONT,size=9,italic=True)
sm.append([])
sm.append(['Source','Items']+cats+['Measurable items','% Covered (of measurable)','% Covered or partly'])
r0=sm.max_row
for s_ in sources:
    rr=sm.max_row+1
    sm.append([s_,f"=COUNTIFS(MASTER!$B:$B,$A{rr})"]+[f"=COUNTIFS(MASTER!$B:$B,$A{rr},MASTER!$H:$H,{get_column_letter(3+i)}${r0})" for i in range(len(cats))]+[f"=C{rr}+D{rr}+E{rr}+F{rr}",f"=IF(J{rr}=0,0,C{rr}/J{rr})",f"=IF(J{rr}=0,0,(C{rr}+D{rr})/J{rr})"])
rt=sm.max_row+1
sm.append(['TOTAL']+[f"=SUM({get_column_letter(c)}{r0+1}:{get_column_letter(c)}{rt-1})" for c in range(2,11)]+[f"=IF(J{rt}=0,0,C{rt}/J{rt})",f"=IF(J{rt}=0,0,(C{rt}+D{rt})/J{rt})"])
for c in sm[r0]: c.font=Font(name=FONT,bold=True,color='FFFFFF',size=10); c.fill=hfill; c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
sm.row_dimensions[r0].height=34
for row in sm.iter_rows(min_row=r0+1,max_row=rt):
    for c in row: c.font=Font(name=FONT,size=10,bold=(row[0].value=='TOTAL'))
    row[10].number_format='0%'; row[11].number_format='0%'
sm.column_dimensions['A'].width=22
for i in range(2,14): sm.column_dimensions[get_column_letter(i)].width=14
# action status table
a0=rt+3
sm.cell(a0,1,'Work still to do (Action status)').font=Font(name=FONT,bold=True,size=12)
sm.append([]); sm.append(['Source','To do','Needs your decision','Done','Not needed'])
h2=sm.max_row
for c in sm[h2]: c.font=Font(name=FONT,bold=True,color='FFFFFF',size=10); c.fill=hfill; c.alignment=Alignment(horizontal='center')
for s_ in sources:
    rr=sm.max_row+1
    sm.append([s_]+[f'=COUNTIFS(MASTER!$B:$B,$A{rr},MASTER!$O:$O,{get_column_letter(2+i)}${h2})' for i in range(4)])
rr=sm.max_row+1
sm.append(['TOTAL']+[f'=SUM({get_column_letter(2+i)}{h2+1}:{get_column_letter(2+i)}{rr-1})' for i in range(4)])
for row in sm.iter_rows(min_row=h2+1,max_row=rr):
    for c in row: c.font=Font(name=FONT,size=10,bold=(row[0].value=='TOTAL'))

# ---- Legend ----
lg=wb.create_sheet('LEGEND')
lines=[('How this sheet works',''),
('One row per item','Every file in the course repo, every page of the LangChain, MCP and Pydantic documentation, every Meridian slide chapter, every Craft page and attachment, every instructor site and other repo.'),
('Read status','What I did with the item itself: Read properly / Read in full, Looked through, NOT READ, Ignored, Excluded by you, Skipped on purpose.'),
('Concept covered?','Whether the concepts in the item are in the books. This is separate from Read status: a page can be unread but its concepts may still be in the books from class notes, and a page can be read but only partly covered.'),
('   Covered','At least 55% of the item\'s distinctive terms are in the relevant book (35% if a chapter links to it as a source).'),
('   Partly covered','25% to 55% of its distinctive terms are in the books, or a chapter links to it. "Missing concepts" lists terms not found.'),
('   Not covered','Under 25% of its distinctive terms are in the books and no chapter links to it.'),
('   No teaching content','Lock files, config, images, job ads, or too little text to measure.'),
('   Excluded by you','Frontend material and the pages you asked me to drop.'),
('   Not applicable','Pages outside the LangChain book (LangGraph not started, other products, integrations, generated API reference).'),
('Overlap score','Share of the item\'s top 40 distinctive terms found in the books. Terms are code names (functions, classes, settings) and heading phrases. Generic words are removed. The books searched depend on the topic: MCP pages against the MCP chapters (LangChain ch11 to ch18), Pydantic pages against the Pydantic book, LangChain pages against the LangChain book, Meridian against the RAG book and chapter 20.'),
('Method','Measured = the term overlap above. Recorded = taken from my build notes (Craft, sites). Proxy = the repo copy of the same file was measured.'),
('Limits of the measure','It finds missing names and headings, not wrong explanations. A high score means the concept is named and discussed, not that the explanation is complete. Code files score low when the book explains the idea in plain words without using the file\'s function names. Treat "Partly covered" as a prompt to look, not a verdict.'),
('Action status','To do = I should add content. Needs your decision = low priority or a choice that is yours (for example community pages). Done = covered. Not needed = nothing to cover.'),
('Updating','After every prompt run: python3 -I build/master_analyze.py <repo> <docsrc>; python3 -I build/master_xlsx.py. Manual updates go in build/master_overrides.json.'),
('Sources of the documentation text','MCP: github.com/modelcontextprotocol/modelcontextprotocol (docs, version 2026-07-28). Pydantic: github.com/pydantic/pydantic (docs folder, main branch, close to but not exactly the "latest" site). LangChain: github.com/langchain-ai/docs (src/oss). The page lists come from the sites\' llms.txt indexes.')]
for a,b in lines:
    lg.append([a,b])
lg.column_dimensions['A'].width=26; lg.column_dimensions['B'].width=120
for row in lg.iter_rows():
    row[0].font=Font(name=FONT,bold=True,size=10); row[1].font=Font(name=FONT,size=10); row[1].alignment=Alignment(wrap_text=True,vertical='top'); row[0].alignment=Alignment(vertical='top')
lg['A1'].font=Font(name=FONT,bold=True,size=14)
wb.move_sheet('SUMMARY',offset=-1)
out=f'{ROOT}/MASTER_COVERAGE.xlsx'; wb.save(out); print('saved',out,len(rows),'rows')
